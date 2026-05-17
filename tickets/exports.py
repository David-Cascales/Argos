import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from django.http import HttpResponse
from tickets.repositories import TicketRepository


class ExportadorTicketsExcel:
    """
    Clase responsable de exportar tickets a Excel.
    """

    CABECERA_COLOR = '2563EB'
    CABECERA_FUENTE_COLOR = 'FFFFFF'

    @staticmethod
    def exportar(queryset=None) -> HttpResponse:
        """
        Genera un archivo Excel con los tickets proporcionados y lo devuelve como respuesta
        HTTP descargable.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Tickets'

        # Estilos de cabecera
        fuente_cabecera = Font(
            bold=True,
            color=ExportadorTicketsExcel.CABECERA_FUENTE_COLOR
        )
        relleno_cabecera = PatternFill(
            start_color=ExportadorTicketsExcel.CABECERA_COLOR,
            end_color=ExportadorTicketsExcel.CABECERA_COLOR,
            fill_type='solid'
        )
        alineacion_centro = Alignment(horizontal='center')

        # Cabeceras
        cabeceras = [
            'ID', 'Título', 'Descripción', 'Estado',
            'Prioridad', 'Categoría', 'Creado por',
            'Fecha creación', 'Fecha actualización', 'Fecha cierre'
        ]
        for col, cabecera in enumerate(cabeceras, start=1):
            celda = ws.cell(row=1, column=col, value=cabecera)
            celda.font = fuente_cabecera
            celda.fill = relleno_cabecera
            celda.alignment = alineacion_centro

        # Datos
        tickets = queryset if queryset is not None else TicketRepository.obtener_todos()
        for fila, ticket in enumerate(tickets, start=2):
            ws.cell(row=fila, column=1, value=ticket.id)
            ws.cell(row=fila, column=2, value=ticket.titulo)
            ws.cell(row=fila, column=3, value=ticket.descripcion)
            ws.cell(row=fila, column=4, value=ticket.get_estado_display())
            ws.cell(row=fila, column=5, value=ticket.get_prioridad_display())
            ws.cell(row=fila, column=6, value=ticket.get_categoria_display())
            ws.cell(row=fila, column=7, value=ticket.creado_por.nombre)
            ws.cell(row=fila, column=8, value=ticket.fecha_creacion.strftime('%d/%m/%Y %H:%M'))
            ws.cell(row=fila, column=9, value=ticket.fecha_actualizacion.strftime('%d/%m/%Y %H:%M'))
            ws.cell(
                row=fila, column=10,
                value=ticket.fecha_cierre.strftime('%d/%m/%Y %H:%M') if ticket.fecha_cierre else ''
            )

        # Ajustar ancho de columnas
        for col in ws.columns:
            max_ancho = 0
            for celda in col:
                if celda.value:
                    max_ancho = max(max_ancho, len(str(celda.value)))
            ws.column_dimensions[col[0].column_letter].width = max_ancho + 4

        # Preparar respuesta HTTP
        response = HttpResponse(
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="tickets.xlsx"'
        wb.save(response)
        return response